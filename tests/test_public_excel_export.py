from datetime import datetime, timezone
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.models.orm import (
    Base, Company, DecisionRun, Job, LinkedInProfileLink,
    JobPostingVersion, ProfileDiscoveryAuthorization, ProfiledSearchOutcome,
    PublicSelectionScope, PublicSelectionScopeItem, PublicTailoredResume, TailoredResume,
)
from app.reporting.public_export import cumulative_workbook_bytes, export_public_workbook


def _session():
    engine = create_engine("sqlite://", future=True)
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine, future=True)()


def _run(session):
    session.add(DecisionRun(
        id="run-1", since_at=datetime(2026, 9, 1, tzinfo=timezone.utc),
        cutoff_at=datetime(2026, 9, 2, tzinfo=timezone.utc), input_timezone="UTC",
        state="collected", policy_snapshot={}, target_count=0,
    ))
    session.flush()


def test_exports_local_run_data_bounded_to_selection_scope(tmp_path):
    session = _session()
    _run(session)
    company = Company(name="Example Co")
    kept = Job(
        source="fixture", external_job_id="kept", company=company,
        company_name_raw="Example Co", title="=Data Analyst", location_raw="Bengaluru",
        is_remote=True, job_url="https://jobs.example/kept",
        apply_url=" https://jobs.example/kept/apply ",
    )
    session.add_all([company, kept])
    session.flush()
    selected_version = JobPostingVersion(
        job_id=kept.id, canonical_duplicate_root_id=kept.id,
        material_content_hash="kept", posting_instance_key="fixture:kept", snapshot={},
    )
    unselected = Job(
        source="fixture", external_job_id="unselected", company=company,
        company_name_raw="Example Co", title="Unselected Role",
    )
    session.add_all([selected_version, unselected])
    session.flush()
    session.add(JobPostingVersion(
        job_id=unselected.id, canonical_duplicate_root_id=unselected.id,
        material_content_hash="unselected", posting_instance_key="fixture:unselected", snapshot={},
    ))
    session.flush()
    session.add_all([
        ProfiledSearchOutcome(
            run_id="run-1", source="fixture", external_job_id="kept", outcome="kept",
            profile_fingerprint="fp", evidence={}, job_id=kept.id,
            posting_version_id=selected_version.id,
        ),
        ProfiledSearchOutcome(
            run_id="run-1", source="fixture", external_job_id="review", outcome="needs_review",
            first_failed_axis="location", reason="remote country is unclear", profile_fingerprint="fp",
            job_url="https://jobs.example/review",
            evidence={"title": "Operations Analyst", "company_name_raw": "Review Co", "location_raw": "Remote"},
        ),
        TailoredResume(job_id=kept.id, score=87.5, artifact_dir="/private/resumes/kept"),
        TailoredResume(job_id=unselected.id, score=10, artifact_dir="/private/resumes/unselected"),
    ])
    session.flush()
    scope = PublicSelectionScope(
        run_id="run-1", revision=1, mode="manual", fingerprint="scope", job_count=1, company_count=1,
    )
    session.add(scope)
    session.flush()
    session.add(PublicSelectionScopeItem(
        scope_id=scope.id, posting_version_id=selected_version.id, company_id=company.id,
    ))
    session.add(PublicTailoredResume(
        posting_version_id=selected_version.id, material_content_hash="kept",
        posting_snapshot={}, score=91.5, artifact_dir="/private/public-resumes/kept",
    ))
    authorization = ProfileDiscoveryAuthorization(
        run_id="run-1", scope_id=scope.id, plan_fingerprint="plan", query_plan=[],
        maximum_calls=1, retry_allowance=0, planned_query_count=1,
    )
    session.add(authorization)
    session.flush()
    session.add(LinkedInProfileLink(
        authorization_id=authorization.id, company_id=company.id,
        linkedin_url="https://www.linkedin.com/in/example", full_name="Example Person",
        headline="Data leader", search_group="data", evidence={"source": "public search"},
    ))
    session.commit()

    output = export_public_workbook(session, "run-1", scope.id, tmp_path / "nested" / "run.xlsx")
    workbook = load_workbook(output, data_only=False)

    assert workbook.sheetnames == [
        "Start here", "All collected", "Jobs", "Companies", "Needs review", "Shortlist",
        "LinkedIn profile links", "Resume artifacts",
    ]
    assert "Every selected job starts as Approved" in workbook["Start here"]["A3"].value
    assert list(workbook["All collected"].values)[1][0:5] == (
        "fixture", "kept", "'=Data Analyst", "Example Co", "Bengaluru",
    )
    assert list(workbook["All collected"].values)[1][7] == "kept"
    assert workbook["All collected"]["K2"].hyperlink.target == "https://jobs.example/kept"
    assert workbook["All collected"]["L2"].hyperlink.target == (
        "https://jobs.example/kept/apply"
    )
    assert list(workbook["Jobs"].values)[1][:6] == (
        True, "fixture", "'=Data Analyst", "Example Co", "Bengaluru", "Remote",
    )
    assert list(workbook["Companies"].values)[1][:4] == (
        True, "Example Co", "employer", None,
    )
    assert list(workbook["Needs review"].values)[1][:6] == (
        "fixture", "review", "Operations Analyst", "Review Co", "Remote", "remote country is unclear",
    )
    assert list(workbook["LinkedIn profile links"].values)[1] == (
        "Example Co", "Example Person", "Data leader", "data",
        "https://www.linkedin.com/in/example", '{"source": "public search"}',
    )
    assert workbook["LinkedIn profile links"]["E2"].hyperlink.target == (
        "https://www.linkedin.com/in/example"
    )
    assert list(workbook["Shortlist"].values)[1][:8] == (
        "Approved", selected_version.id, company.id, "Example Co", kept.id,
        "fixture", "kept", "'=Data Analyst",
    )
    assert workbook["Shortlist"]["I2"].hyperlink.target == "https://jobs.example/kept"
    assert str(workbook["Shortlist"].data_validations.dataValidation[0].sqref) == "A2"
    assert workbook["Shortlist"].data_validations.dataValidation[0].formula1 == (
        '"Approved,Declined"'
    )
    assert list(workbook["Resume artifacts"].values)[1][1:7] == (
        selected_version.id, kept.id, "'=Data Analyst", "Example Co", 91.5,
        "/private/public-resumes/kept",
    )
    assert workbook["Resume artifacts"].max_row == 2
    assert "/private/resumes/kept" not in str(list(workbook["Resume artifacts"].values))

    cumulative = load_workbook(BytesIO(cumulative_workbook_bytes(session, ["run-1"])))
    assert cumulative.sheetnames == ["All collected", "Eligible jobs", "Companies"]
    assert cumulative["Eligible jobs"].max_row == 2
    eligible_headers = [cell.value for cell in cumulative["Eligible jobs"][1]]
    job_url_column = eligible_headers.index("job_url") + 1
    assert cumulative["Eligible jobs"].cell(2, job_url_column).hyperlink.target == (
        "https://jobs.example/kept"
    )


def test_export_includes_empty_sheets_with_headers(tmp_path):
    session = _session()
    _run(session)

    scope = PublicSelectionScope(
        run_id="run-1", revision=1, mode="manual", fingerprint="empty", job_count=0, company_count=0,
    )
    session.add(scope)
    session.commit()

    output = export_public_workbook(session, "run-1", scope.id, tmp_path / "empty.xlsx")
    workbook = load_workbook(output)

    assert workbook["Start here"].max_row == 6
    assert all(
        sheet.max_row == 1
        for sheet in workbook.worksheets
        if sheet.title != "Start here"
    )


def test_export_rejects_a_scope_from_another_run(tmp_path):
    session = _session()
    _run(session)
    session.add(DecisionRun(
        id="run-2", since_at=datetime(2026, 9, 1, tzinfo=timezone.utc),
        cutoff_at=datetime(2026, 9, 2, tzinfo=timezone.utc), input_timezone="UTC",
        state="collected", policy_snapshot={}, target_count=0,
    ))
    session.flush()
    scope = PublicSelectionScope(
        run_id="run-2", revision=1, mode="manual", fingerprint="other", job_count=0, company_count=0,
    )
    session.add(scope)
    session.commit()

    try:
        export_public_workbook(session, "run-1", scope.id, tmp_path / "wrong.xlsx")
    except ValueError as exc:
        assert "selection scope" in str(exc)
    else:
        raise AssertionError("scope from another run was accepted")


def test_export_rejects_phase_b_research_scope(tmp_path):
    session = _session()
    _run(session)
    scope = PublicSelectionScope(
        run_id="run-1", revision=1, mode="all_eligible", purpose="phase_b_research",
        fingerprint="research", job_count=0, company_count=0,
    )
    session.add(scope); session.commit()

    with pytest.raises(ValueError, match="final selection scope"):
        export_public_workbook(session, "run-1", scope.id, tmp_path / "research.xlsx")


def test_export_hides_links_for_companies_excluded_from_a_later_scope(tmp_path):
    session = _session()
    _run(session)
    included = Company(name="Included Co")
    excluded = Company(name="Excluded Co")
    session.add_all([included, excluded])
    session.flush()
    first = PublicSelectionScope(
        run_id="run-1", revision=1, mode="manual", fingerprint="first", job_count=2, company_count=2,
    )
    latest = PublicSelectionScope(
        run_id="run-1", revision=2, mode="manual", fingerprint="latest", job_count=1, company_count=1,
    )
    session.add_all([first, latest])
    session.flush()
    session.add_all([
        PublicSelectionScopeItem(scope_id=first.id, posting_version_id=101, company_id=included.id),
        PublicSelectionScopeItem(scope_id=first.id, posting_version_id=102, company_id=excluded.id),
        PublicSelectionScopeItem(scope_id=latest.id, posting_version_id=103, company_id=included.id),
    ])
    authorization = ProfileDiscoveryAuthorization(
        run_id="run-1", scope_id=first.id, plan_fingerprint="first-plan", query_plan=[],
        maximum_calls=0, retry_allowance=0, planned_query_count=0,
    )
    session.add(authorization)
    session.flush()
    session.add_all([
        LinkedInProfileLink(
            authorization_id=authorization.id, company_id=included.id,
            linkedin_url="https://linkedin.com/in/included", full_name="Included", headline=None,
            search_group="data", evidence={},
        ),
        LinkedInProfileLink(
            authorization_id=authorization.id, company_id=excluded.id,
            linkedin_url="https://linkedin.com/in/excluded", full_name="Excluded", headline=None,
            search_group="data", evidence={},
        ),
    ])
    session.commit()

    output = export_public_workbook(session, "run-1", latest.id, tmp_path / "latest.xlsx")
    values = list(load_workbook(output)["LinkedIn profile links"].values)

    assert values == [
        ("company", "name", "headline", "search_group", "linkedin_url", "evidence"),
        ("Included Co", "Included", None, "data", "https://linkedin.com/in/included", "{}"),
    ]
