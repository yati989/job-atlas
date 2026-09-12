"""Offline golden-path verification for the attended Decision Run workflow.

Internal persistence, ordering, attachment and lifecycle seams are real. The
external Gmail boundary is an ``InMemoryMailbox`` and the external resume
renderer is a deterministic PDF fixture; neither sends email nor invokes TeX.
"""
from argparse import Namespace
from contextlib import nullcontext
from datetime import datetime, timedelta, timezone

from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.dashboard import queries
from app.decision_runs import (
    ChosenPairing,
    CompanyEnrichmentResult,
    CompanyProcessingOutcome,
    ContactEnrichmentResult,
    ContactOutcome,
    EnrichmentResult,
    GmailOutcome,
    GmailResult,
    PhaseBEvidenceOutcome,
    TailoringOutcome,
    TailoringResult,
    approve_decision_run,
    create_decision_run,
    finalize_decision_run,
    freeze_contact_enrichment_manifest,
    freeze_draft_preparation_manifest,
    freeze_gmail_draft_manifest,
    freeze_job_enrichment_manifest,
    freeze_resume_tailoring_manifest,
    report_company_phase_a,
    report_company_phase_b,
    report_contact_enrichment,
    report_draft_preparation,
    report_final_report,
    report_job_enrichment,
    report_resume_tailoring,
    screen_decision_run,
)
from app.decision_runs.ingestion_funnel import report_canonical_jobs, report_ingestion_funnel
from app.decision_runs.outreach_funnel import DraftOutcome, DraftResult, begin_final_report
from app.decision_runs.progress import (StageCounts, StageName, finish_stage, interrupt_stage,
                                        read_run_progress, start_stage)
from app.decision_runs.ranking import RankableJob, build_ranking_plan, group_for
from app.models.orm import Base, Contact, DecisionRun, OutreachDeliveryAttempt, OutreachMessage, TailoredResume
from app.models.schemas import NormalizedJob
from app.outreach import cli as outreach_cli
from app.outreach import drafts
from app.outreach.state import InMemoryMailbox
from app.pipeline.runner import FetchObservation, GateOutcome, IngestionGateFacts
from app.pipeline.upsert import upsert_job
from app.reporting.final_run_report import build_final_workbook
from app.resume.render import RenderResult
from app.resume.schema import EXAMPLE_MASTER_PATH, load_master
from app.resume.tailor import save_tailored


NOW = datetime(2026, 8, 25, 12, tzinfo=timezone.utc)


def _stage(progress, name):
    return next(item for item in progress.stages if item.name == name.value)


def test_offline_attended_pipeline_uses_one_immutable_scope_and_only_creates_a_gmail_draft(tmp_path, monkeypatch):
    engine = create_engine("sqlite:///:memory:", future=True)
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine, future=True)()
    run = create_decision_run(session, since=NOW - timedelta(days=1), cutoff=NOW)

    kept = NormalizedJob(
        source="fixture", external_job_id="kept", title="Senior Data Engineer",
        company_name_raw="Acme", description_raw="Build a remote data platform.",
        posted_at=NOW - timedelta(hours=2), is_remote=True,
        job_url="https://jobs.example/kept", apply_url="https://apply.example/kept",
    )
    rejected = NormalizedJob(
        source="fixture", external_job_id="rejected", title="Account Executive",
        company_name_raw="Other", description_raw="Sell software.", posted_at=NOW,
        is_remote=True, job_url="https://jobs.example/rejected",
    )
    job = upsert_job(session, kept)

    # Fetch telemetry is fed into the production ingestion funnel; no fake
    # lifecycle object or global pending queue is used.
    start_stage(session, run.id, StageName.FETCH_JOBS, expected_count=2, reason="fixture fetch")
    finish_stage(session, run.id, StageName.FETCH_JOBS, counts=StageCounts(2, 2, 0, 0, 0))
    kept_observation = FetchObservation(kept, 1, 1)
    rejected_observation = FetchObservation(rejected, 1, 2)
    kept_fact = GateOutcome(kept_observation, 1, "advanced", connector_dimensions={"country": "India"})
    rejected_fact = GateOutcome(rejected_observation, 1, "dropped", "role_mismatch")
    canonical = report_ingestion_funnel(session, run.id, IngestionGateFacts(
        raw=(kept_fact, rejected_fact),
        collection=(kept_fact, rejected_fact),
        relevance=(kept_fact, rejected_fact),
    ))
    report_canonical_jobs(session, run.id, canonical)

    enrichment = freeze_job_enrichment_manifest(session, run.id)
    assert enrichment.posting_version_ids and len(enrichment.posting_version_ids) == 1
    report_job_enrichment(session, run.id, [EnrichmentResult(enrichment.posting_version_ids[0], "enriched")])

    # Screening freezes company scope before company enrichment; ranking waits
    # for both evidence sources to report a terminal outcome.
    screen_decision_run(session, run.id)
    report_company_phase_a(session, run.id, [
        CompanyEnrichmentResult(job.company_id, CompanyProcessingOutcome.ENRICHED),
    ])
    report_company_phase_b(session, run.id, [
        CompanyEnrichmentResult(job.company_id, CompanyProcessingOutcome.ENRICHED,
                                evidence_outcome=PhaseBEvidenceOutcome.CONFIRMED_MISSING,
                                source="glassdoor", evidence_url="https://glassdoor.example/acme"),
        CompanyEnrichmentResult(job.company_id, CompanyProcessingOutcome.ENRICHED,
                                evidence_outcome=PhaseBEvidenceOutcome.CONFIRMED_MISSING,
                                source="ambitionbox", evidence_url="https://ambitionbox.example/acme"),
    ])
    finalize_decision_run(session, run.id)
    waiting = _stage(read_run_progress(session, run.id), StageName.APPROVAL_GATE)
    assert waiting.status == "waiting_for_approval"

    scope = approve_decision_run(
        session, run.id, "target 1; G1 all; G2 all; G3 all; G4 all",
        approver="offline-e2e", contact_search_selections={job.company_id: "data_ai"},
    )
    assert scope.posting_version_ids == enrichment.posting_version_ids

    # Mutating the current job after approval must not expand or replace the
    # frozen posting-version scope used by the downstream stages.
    job.description_raw = "later mutable description"
    resume_manifest = freeze_resume_tailoring_manifest(session, run.id)
    version_id = resume_manifest.posting_version_ids[0]
    interrupt_stage(session, run.id, StageName.RESUME_TAILORING,
                    reason="offline worker interruption")
    resumed = freeze_resume_tailoring_manifest(session, run.id)
    assert resumed.posting_version_ids == resume_manifest.posting_version_ids
    resume_progress = _stage(read_run_progress(session, run.id), StageName.RESUME_TAILORING)
    assert [(attempt.attempt_number, attempt.status) for attempt in resume_progress.attempts] == [
        (1, "interrupted"), (2, "running"),
    ]
    def render_fixture(_master, artifact_dir, *, basename):
        artifact_dir.mkdir(parents=True, exist_ok=True)
        resume_path = artifact_dir / f"{basename}.pdf"
        resume_path.write_bytes(b"%PDF-1.4 offline fixture resume")
        return RenderResult(artifact_dir / f"{basename}.tex", resume_path, ok=True)

    monkeypatch.setattr("app.resume.tailor.render_resume", render_fixture)
    gap = {"score": 80, "must_haves": [], "nice_to_haves": [],
           "surfaceable_gaps": [], "real_gaps": []}
    tailored, rendered = save_tailored(
        session, job_id=job.id, tailored=load_master(EXAMPLE_MASTER_PATH), score=80, gap_report=gap,
        out_root=tmp_path / "tailored", decision_run_id=run.id,
    )
    resume_path = rendered.pdf_path
    assert tailored.artifact_dir == str(resume_path.parent)
    assert resume_path.is_file()
    assert _stage(read_run_progress(session, run.id), StageName.RESUME_TAILORING).status == "completed"

    contact = Contact(company_id=job.company_id, full_name="Jane Doe", title="Data Engineer",
                      linkedin_url="https://linkedin.example/jane", email_guess="jane@acme.example",
                      seniority_tier="head")
    session.add(contact)
    session.flush()
    contact_manifest = freeze_contact_enrichment_manifest(session, run.id)
    assert [(row.company_id, row.search_group) for row in contact_manifest] == [(job.company_id, "data_ai")]
    report_contact_enrichment(session, run.id, [ContactEnrichmentResult(
        job.company_id, "data_ai", ContactOutcome.ENRICHED, new_contact_ids=(contact.id,),
        coverage={"new_contact_evidence": [{"contact_id": contact.id, "search_group": "data_ai",
                                              "evidence": "Current Head of Data"}]},
    )])

    draft_manifest = freeze_draft_preparation_manifest(session, run.id, chosen_pairings=[
        ChosenPairing(contact.id, job.id, version_id),
    ])
    target = drafts.resolve_recipient(session, to_email=contact.email_guess)
    draft = drafts.save_draft(
        session, target, to_name="Jane Doe", recipient_title="Head of Data", recipient_tier="head",
        resume_kind="tailored", pairing_reason=draft_manifest[0].pairing_reason,
        subject="Data platform opportunity", body="Hello Jane", tailored_resume_id=tailored.id,
        resume_path=str(resume_path), matched_job_id=job.id, decision_run_id=run.id,
        posting_version_id=version_id,
    )
    report_draft_preparation(session, run.id, [DraftResult(contact.id, DraftOutcome.PREPARED, outreach_draft_id=draft.id)])

    gmail_manifest = freeze_gmail_draft_manifest(session, run.id)
    mailbox = InMemoryMailbox()
    selected = []
    from app.outreach import state as outreach_state
    actual_push = outreach_state.push_attempt
    def record_push(active_session, active_mailbox, attempt):
        selected.append(attempt.id)
        return actual_push(active_session, active_mailbox, attempt)
    monkeypatch.setattr(outreach_cli, "get_session", lambda: nullcontext(session))
    monkeypatch.setattr(outreach_cli.gmail_ops, "get_service", lambda: object())
    monkeypatch.setattr(outreach_cli.gmail_ops, "GmailMailboxAdapter", lambda _service: mailbox)
    monkeypatch.setattr(outreach_state, "push_attempt", record_push)
    # Exercise the production run-scoped selector: it selects every frozen,
    # held approved attempt and writes Gmail Drafts rather than sending.
    outreach_cli._cmd_push(Namespace(run_id=run.id, all=False, to_email=None))
    assert len(mailbox.drafts) == 1 and mailbox.sent == {}
    assert next(iter(mailbox.drafts.values()))["attachment_path"] == str(resume_path)
    assert selected == [gmail_manifest[0].delivery_attempt_id]
    assert session.get(OutreachMessage, session.get(OutreachDeliveryAttempt, selected[0]).message_id).is_calibration is False

    final_path = tmp_path / "final.xlsx"
    assert build_final_workbook(session, run.id, final_path) == final_path
    workbook = load_workbook(final_path, read_only=True)
    try:
        assert {"Fresh outreach companies", "Jobs and resumes", "Delivery state", "Run summary"} <= set(workbook.sheetnames)
    finally:
        workbook.close()
    begin_final_report(session, run.id)
    report_final_report(session, run.id, output=str(final_path),
                        self_report_requested=False, self_report_delivered=False)

    progress = read_run_progress(session, run.id)
    statuses = {item.name: item.status for item in progress.stages}
    assert all(statuses[name.value] == "completed" for name in (
        StageName.FETCH_JOBS, StageName.COLLECTION_DEDUPLICATION, StageName.RELEVANCE_STORAGE,
        StageName.JOB_DEDUPLICATION, StageName.JOB_ENRICHMENT, StageName.COMPANY_DEDUPLICATION,
        StageName.COMPANY_PHASE_A, StageName.COMPANY_PHASE_B, StageName.SCREENING_RANKING_GROUPING,
        StageName.APPROVAL_GATE, StageName.RESUME_TAILORING, StageName.CONTACT_ENRICHMENT,
        StageName.DRAFT_PREPARATION, StageName.GMAIL_DRAFT_POSTING, StageName.FINAL_REPORT,
    ))
    assert _stage(progress, StageName.APPROVAL_GATE).selected_company_count == 1

    # Dashboard reads use the same durable records and preserve exact reasons
    # and snapshot links rather than rereading the mutable job row.
    job_funnel = queries.decision_run_job_funnel(session, run.id)
    assert list(job_funnel.loc[job_funnel.stage == "relevance_storage", "dropped"])[0] == 1
    assert queries.decision_run_job_funnel_drilldown(session, run.id, stage_name="relevance_storage").iloc[1].reason == "role_mismatch"
    assert queries.decision_run_company_funnel(session, run.id).iloc[0].companies == 1
    assert queries.decision_run_resume_tailoring_drilldown(session, run.id).iloc[0].job_url == kept.job_url
    assert queries.decision_run_contact_funnel_drilldown(session, run.id).iloc[0].outcome == "enriched"
    assert queries.decision_run_outreach_funnel_drilldown(session, run.id, stage_name="gmail_draft_posting").iloc[0].outcome == GmailOutcome.POSTED.value


def test_dashboard_progress_isolated_for_two_concurrent_active_runs():
    """The selector/read model must never blend concurrent Decision Run state."""
    engine = create_engine("sqlite:///:memory:", future=True)
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine, future=True)()
    for run_id in ("run-a", "run-b"):
        session.add(DecisionRun(id=run_id, since_at=NOW - timedelta(days=1), cutoff_at=NOW,
                                input_timezone="UTC", state="preparing", policy_snapshot={},
                                target_count=1, telemetry_version="decision-run-progress-v1"))
    session.flush()
    start_stage(session, "run-a", StageName.FETCH_JOBS, expected_count=1, reason="run a collecting")
    start_stage(session, "run-b", StageName.FETCH_JOBS, expected_count=2, reason="run b collecting")
    finish_stage(session, "run-b", StageName.FETCH_JOBS, counts=StageCounts(2, 2, 0, 0, 0))

    a = queries.decision_run_progress(session, "run-a")
    b = queries.decision_run_progress(session, "run-b")
    assert a.iloc[0][["status", "expected", "pending"]].tolist() == ["running", 1, 1]
    assert b.iloc[0][["status", "expected", "advanced"]].tolist() == ["completed", 2, 2]
    options = queries.decision_run_options(session)
    assert set(options.run_id) == {"run-a", "run-b"}


def test_real_ranking_policy_emits_all_four_company_groups():
    """The terminal company dashboard's group labels come from one ranking policy."""
    jobs = (
        RankableJob(1, 1, 1, "data_engineering", None, None, NOW, True),   # remote, salary unknown
        RankableJob(2, 2, 2, "data_engineering", 25, None, NOW, False),    # high salary, not remote
        RankableJob(3, 3, 3, "data_engineering", 10, None, NOW, True),     # remote, lower salary
        RankableJob(4, 4, 4, "data_engineering", 10, None, NOW, False),    # lower salary, not remote
    )
    assert [group_for(job.is_remote, job.effective_salary_lpa) for job in jobs] == [1, 2, 3, 4]
    plan = build_ranking_plan(jobs)
    assert [(item.company_id, item.group_number) for item in plan.companies] == [(1, 1), (2, 2), (3, 3), (4, 4)]
