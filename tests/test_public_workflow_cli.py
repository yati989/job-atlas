from datetime import datetime, timezone
import json
from shutil import copyfile
from openpyxl import load_workbook

from app.models.orm import (
    Company, DecisionRun, GuidedSourceRun, Job, JobPostingVersion,
    ProfiledSearchOutcome, PublicWorkflowStage,
)
from app.workflows.public_database import public_session
from job_search_agent import main


def _seed(private_home):
    with public_session(private_home) as session:
        session.add(DecisionRun(
            id="run-1", since_at=datetime.now(timezone.utc),
            cutoff_at=datetime.now(timezone.utc), input_timezone="UTC",
            state="collected", policy_snapshot={
                "version": 1, "name": "design", "professions": ["product design"],
                "search_terms": ["product designer"],
                "relevance_terms": ["product designer"], "countries": ["IN"],
                "arrangements": ["remote"], "seniority": ["individual_contributor"],
                "collection_window_days": 30,
                "source_selection": {"mode": "all_compatible"},
            }, target_count=0,
        ))
        session.add(GuidedSourceRun(
            run_id="run-1", source="fixture", status="completed",
            outcome_counts={"kept": 1},
        ))
        company = Company(name="DesignCo")
        session.add(company); session.flush()
        job = Job(
            source="fixture", external_job_id="1", company_id=company.id,
            company_name_raw=company.name, title="Product Designer",
            location_raw="Remote - India", is_remote=True,
            description_raw="Requires Figma.", enrichment_status="pending",
        )
        session.add(job); session.flush()
        version = JobPostingVersion(
            job_id=job.id, canonical_duplicate_root_id=job.id,
            material_content_hash="hash-1", posting_instance_key="fixture:1",
            snapshot={"description_raw": job.description_raw},
        )
        session.add(version); session.flush()
        session.add(ProfiledSearchOutcome(
            run_id="run-1", source="fixture", external_job_id="1", outcome="kept",
            profile_fingerprint="fp", evidence={}, job_id=job.id,
            posting_version_id=version.id,
        ))
        session.flush()
        return company.id, version.id


def test_public_phase_selection_tailoring_and_export_commands_share_private_sqlite(
    tmp_path, capsys,
):
    private_home = tmp_path / "private"
    company_id, posting_version_id = _seed(private_home)

    assert main([
        "phase-a-context", "--run-id", "run-1",
        "--private-home", str(private_home),
    ]) == 0
    context = json.loads(capsys.readouterr().out)
    assert context["jobs"][0]["posting_version_id"] == posting_version_id

    results = tmp_path / "phase-a.json"
    results.write_text(json.dumps({
        "context_fingerprint": context["context_fingerprint"],
        "jobs": [{
            "posting_version_id": posting_version_id,
            "status": "done", "hard_skills": ["Figma"],
        }],
        "companies": [{
            "company_id": company_id, "company_type": "employer",
            "search_groups": ["product_design"],
        }],
    }), encoding="utf-8")
    assert main([
        "apply-phase-a", "--run-id", "run-1", "--input", str(results),
        "--private-home", str(private_home),
    ]) == 0
    assert json.loads(capsys.readouterr().out)["jobs_recorded"] == 1

    assert main([
        "freeze-selection", "--run-id", "run-1", "--mode", "all_eligible",
        "--phase-b-research", "--confirm", "--private-home", str(private_home),
    ]) == 0
    research_scope = json.loads(capsys.readouterr().out)
    assert research_scope["job_count"] == 1
    assert research_scope["purpose"] == "phase_b_research"

    assert main([
        "phase-b-plan", "--scope-id", str(research_scope["scope_id"]),
        "--source", "glassdoor", "--private-home", str(private_home),
    ]) == 0
    phase_b_plan = json.loads(capsys.readouterr().out)
    assert phase_b_plan["planned_call_count"] == 1
    assert main([
        "authorize-phase-b", "--scope-id", str(research_scope["scope_id"]),
        "--source", "glassdoor", "--maximum-calls", "1", "--confirm",
        "--private-home", str(private_home),
    ]) == 0
    phase_b_auth = json.loads(capsys.readouterr().out)
    assert main([
        "reserve-phase-b", "--authorization-id", str(phase_b_auth["authorization_id"]),
        "--company-id", str(company_id), "--source", "glassdoor", "--confirm",
        "--private-home", str(private_home),
    ]) == 0
    phase_b_call = json.loads(capsys.readouterr().out)
    phase_b_result = tmp_path / "phase-b.json"
    phase_b_result.write_text(json.dumps({
        "company_id": company_id, "source": "glassdoor", "status": "missing",
        "evidence": {"reason": "no verified match"},
    }), encoding="utf-8")
    assert main([
        "apply-phase-b", "--call-id", str(phase_b_call["call_id"]),
        "--input", str(phase_b_result), "--private-home", str(private_home),
    ]) == 0
    assert json.loads(capsys.readouterr().out)["stage_status"] == "completed"

    assert main([
        "status", "--run-id", "run-1", "--private-home", str(private_home),
    ]) == 0
    mid_status = json.loads(capsys.readouterr().out)
    assert {item["name"]: item["status"] for item in mid_status["stages"]}["selection"] == "not_requested"

    filters = tmp_path / "filters.yaml"
    filters.write_text(
        "version: 1\ngroups:\n  - name: designers\n    title_terms: [designer]\n",
        encoding="utf-8",
    )
    assert main([
        "selection-preview", "--run-id", "run-1", "--config", str(filters),
        "--private-home", str(private_home),
    ]) == 0
    assert json.loads(capsys.readouterr().out)["retained_jobs"] == 1

    assert main([
        "freeze-selection", "--run-id", "run-1", "--mode", "filtered",
        "--config", str(filters), "--confirm", "--private-home", str(private_home),
    ]) == 0
    frozen = json.loads(capsys.readouterr().out)
    assert frozen["job_count"] == 1
    assert frozen["revision"] == 2
    assert frozen["purpose"] == "selection"

    assert main([
        "tailoring-context", "--scope-id", str(frozen["scope_id"]),
        "--private-home", str(private_home),
    ]) == 0
    tailoring = json.loads(capsys.readouterr().out)
    assert tailoring["needs_tailoring"][0]["posting_version_id"] == posting_version_id

    workbook = tmp_path / "run.xlsx"
    assert main([
        "export", "--run-id", "run-1", "--scope-id", str(frozen["scope_id"]),
        "--out", str(workbook),
        "--private-home", str(private_home),
    ]) == 0
    assert json.loads(capsys.readouterr().out)["workbook"] == str(workbook)
    assert workbook.exists()

    reviewed_path = tmp_path / "run-reviewed.xlsx"
    copyfile(workbook, reviewed_path)
    reviewed_workbook = load_workbook(reviewed_path)
    reviewed_workbook["Shortlist"]["A2"] = "Approved"
    reviewed_workbook.save(reviewed_path)
    assert main([
        "apply-workbook-shortlist", "--run-id", "run-1",
        "--scope-id", str(frozen["scope_id"]), "--input", str(reviewed_path),
        "--confirm", "--private-home", str(private_home),
    ]) == 0
    workbook_scope = json.loads(capsys.readouterr().out)
    assert workbook_scope["approved_job_count"] == 1
    assert workbook_scope["revision"] == 3
    with public_session(private_home) as session:
        export_stage = session.query(PublicWorkflowStage).filter_by(
            run_id="run-1", name="export",
        ).one()
        assert export_stage.detail["workbook"] == str(reviewed_path)
        assert export_stage.detail["scope_id"] == workbook_scope["scope_id"]

    reviewed_workbook = load_workbook(reviewed_path)
    reviewed_workbook["Shortlist"]["A2"] = "Declined"
    reviewed_workbook.save(reviewed_path)
    assert main([
        "apply-workbook-shortlist", "--run-id", "run-1",
        "--scope-id", str(frozen["scope_id"]), "--input", str(reviewed_path),
        "--confirm", "--private-home", str(private_home),
    ]) == 0
    declined_scope = json.loads(capsys.readouterr().out)
    assert declined_scope["approved_job_count"] == 0
    assert declined_scope["company_count"] == 0

    assert main([
        "status", "--run-id", "run-1", "--private-home", str(private_home),
    ]) == 0
    status = json.loads(capsys.readouterr().out)
    stages = {item["name"]: item["status"] for item in status["stages"]}
    assert stages["phase_a"] == "completed"
    assert stages["phase_b"] == "completed"
    assert stages["selection"] == "completed"
    assert stages["tailoring"] == "running"
    assert stages["export"] == "completed"
