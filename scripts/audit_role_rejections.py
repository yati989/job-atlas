"""Create a diagnostic-only workbook of central-gate role rejections.

This deliberately calls connectors and the shared fetch/retry/window/gate
helpers, but never calls ``gate_and_upsert``: no jobs, progress state, or
completion email are persisted.  It is a review artifact, not a vocabulary
change mechanism.

Usage:
    python -m scripts.audit_role_rejections
"""
from __future__ import annotations

import argparse
import json
import logging
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

from app.pipeline.registry import ACTIVE_CONNECTORS
from app.pipeline.relevance import first_failing_axis
from app.pipeline.runner import FetchResult, fetch_all, retry_failed
from app.pipeline.window import apply_sync_window, parse_since
from app.pipeline.workload import enforce_workload_bounds
from scripts.run_headed_sources import CONNECTORS as HEADED_CONNECTORS


logger = logging.getLogger("audit_role_rejections")
DEFAULT_CUTOFF = "2026-08-21T21:00:00+05:30"
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
TITLE_FILL = PatternFill("solid", fgColor="D9EAF7")
STATUS_PARTIAL_FILL = PatternFill("solid", fgColor="FFF2CC")
STATUS_FAILURE_FILL = PatternFill("solid", fgColor="F4CCCC")
THIN_BLUE = Side(style="thin", color="9EADBA")


def _utc_excel(value: Any) -> datetime | None:
    """Return a timezone-free UTC value, which Excel can store as a date."""
    if not isinstance(value, datetime):
        return None
    if value.tzinfo is None:
        return value
    return value.astimezone(timezone.utc).replace(tzinfo=None)


def _dimensions(connector: object) -> tuple[str | None, str | None, str]:
    search = getattr(connector, "search", None)
    if search is None:
        terms = getattr(connector, "search_terms", None)
        search = ", ".join(terms) if isinstance(terms, (list, tuple)) else None
    location = getattr(connector, "location_mode", None)
    values = {
        name: getattr(connector, name)
        for name in ("search", "search_terms", "location_mode", "country")
        if getattr(connector, name, None) is not None
    }
    return (str(search) if search is not None else None,
            str(location) if location is not None else None,
            json.dumps(values, sort_keys=True))


def _result_status(result: FetchResult) -> tuple[str, str | None, str | None]:
    if result.error is None:
        return "complete", None, None
    if result.jobs:
        return "partial", type(result.error).__name__, str(result.error)
    return "failed", type(result.error).__name__, str(result.error)


def _job_url(job: Any) -> str | None:
    return getattr(job, "job_url", None) or getattr(job, "apply_url", None)


def _audit_rows(results: list[FetchResult], cutoff_at: datetime) -> tuple[list[dict], list[dict]]:
    """Return raw occurrences and source/external-id deduplicated role drops."""
    raw_rows: list[dict] = []
    unique_rows: list[dict] = []
    seen: set[tuple[str, str]] = set()
    for result in results:
        if not result.jobs:
            continue
        search, location_mode, _ = _dimensions(result.connector)
        for job in result.jobs:
            if first_failing_axis(job, cutoff_at=cutoff_at) != "role":
                continue
            row = {
                "source": result.source,
                "search_term": search,
                "location_mode": location_mode,
                "title": getattr(job, "title", None),
                "company": getattr(job, "company_name_raw", None),
                "location": getattr(job, "location_raw", None),
                "is_remote": getattr(job, "is_remote", None),
                "posting_date_utc": _utc_excel(getattr(job, "posted_at", None)),
                "external_id": getattr(job, "external_job_id", None),
                "job_link": _job_url(job),
            }
            raw_rows.append(row)
            key = (result.source, str(row["external_id"]))
            if key not in seen:
                seen.add(key)
                unique_rows.append(row)
    return raw_rows, unique_rows


def _style_sheet(ws, *, widths: dict[int, float], freeze: str, table_name: str) -> None:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    for column, width in widths.items():
        ws.column_dimensions[get_column_letter(column)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    if ws.max_row >= 2:
        table = Table(displayName=table_name, ref=ws.dimensions)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False,
            showLastColumn=False, showRowStripes=True, showColumnStripes=False,
        )
        ws.add_table(table)


def _write_workbook(
    output: Path,
    *,
    raw_rows: list[dict],
    unique_rows: list[dict],
    results: list[FetchResult],
    cutoff_at: datetime,
    started_at: datetime,
) -> None:
    wb = Workbook()
    titles_ws = wb.active
    titles_ws.title = "Unique Titles"
    all_ws = wb.create_sheet("All Rejected Rows")
    coverage_ws = wb.create_sheet("Coverage")

    raw_by_title: dict[str, list[dict]] = defaultdict(list)
    unique_by_title: dict[str, list[dict]] = defaultdict(list)
    for row in raw_rows:
        raw_by_title[row["title"] or "(missing title)"].append(row)
    for row in unique_rows:
        unique_by_title[row["title"] or "(missing title)"].append(row)

    titles_ws.append([
        "Rejected job title", "Raw occurrences", "Unique source jobs",
        "Sources", "Search terms", "Example companies", "Example link",
    ])
    for title in sorted(raw_by_title, key=lambda value: (-len(raw_by_title[value]), value.casefold())):
        raw = raw_by_title[title]
        unique = unique_by_title[title]
        titles_ws.append([
            title, len(raw), len(unique),
            ", ".join(sorted({str(row["source"]) for row in raw})),
            ", ".join(sorted({row["search_term"] for row in raw if row["search_term"]})),
            ", ".join(sorted({row["company"] for row in raw if row["company"]}))[:500],
            next((row["job_link"] for row in raw if row["job_link"]), None),
        ])
    _style_sheet(titles_ws, widths={1: 44, 2: 16, 3: 18, 4: 28, 5: 34, 6: 42, 7: 56}, freeze="A2", table_name="UniqueTitles")
    for row in range(2, titles_ws.max_row + 1):
        cell = titles_ws.cell(row, 7)
        if cell.value:
            cell.hyperlink = cell.value
            cell.style = "Hyperlink"

    all_ws.append([
        "Source", "Search term", "Location mode", "Title", "Company", "Location",
        "Remote", "Posting date (UTC)", "External ID", "Direct job link",
    ])
    for row in sorted(raw_rows, key=lambda item: (item["title"] or "", item["source"], item["external_id"] or "")):
        all_ws.append([
            row["source"], row["search_term"], row["location_mode"], row["title"],
            row["company"], row["location"], row["is_remote"], row["posting_date_utc"],
            row["external_id"], row["job_link"],
        ])
    _style_sheet(all_ws, widths={1: 20, 2: 26, 3: 18, 4: 46, 5: 28, 6: 30, 7: 11, 8: 21, 9: 28, 10: 60}, freeze="A2", table_name="AllRejectedRows")
    for row in range(2, all_ws.max_row + 1):
        all_ws.cell(row, 8).number_format = "yyyy-mm-dd hh:mm"
        cell = all_ws.cell(row, 10)
        if cell.value:
            cell.hyperlink = cell.value
            cell.style = "Hyperlink"

    coverage_ws.append([
        "Source", "Instance", "Search term", "Location mode", "Dimensions",
        "Fetched", "Raw role-rejected", "Unique role-rejected", "Status",
        "Attempts", "Fetch seconds", "Error type", "Error", "Started (UTC)", "Ended (UTC)",
    ])
    for number, result in enumerate(results, 1):
        status, error_type, error = _result_status(result)
        search, location, dimensions = _dimensions(result.connector)
        result_rows = []
        result_unique = set()
        for job in result.jobs or []:
            if first_failing_axis(job, cutoff_at=cutoff_at) == "role":
                result_rows.append(job)
                result_unique.add((job.source, job.external_job_id))
        coverage_ws.append([
            result.source, number, search, location, dimensions, len(result.jobs or []),
            len(result_rows), len(result_unique), status, result.attempts,
            round(result.elapsed_s, 1), error_type, error,
            _utc_excel(result.started_at), _utc_excel(result.ended_at),
        ])
    _style_sheet(coverage_ws, widths={1: 20, 2: 10, 3: 26, 4: 18, 5: 46, 6: 12, 7: 18, 8: 20, 9: 12, 10: 10, 11: 14, 12: 20, 13: 56, 14: 21, 15: 21}, freeze="A2", table_name="Coverage")
    for row in range(2, coverage_ws.max_row + 1):
        coverage_ws.cell(row, 14).number_format = "yyyy-mm-dd hh:mm"
        coverage_ws.cell(row, 15).number_format = "yyyy-mm-dd hh:mm"
    coverage_ws.conditional_formatting.add(
        f"I2:I{coverage_ws.max_row}",
        FormulaRule(formula=["I2=\"partial\""], fill=STATUS_PARTIAL_FILL),
    )
    coverage_ws.conditional_formatting.add(
        f"I2:I{coverage_ws.max_row}",
        FormulaRule(formula=["I2=\"failed\""], fill=STATUS_FAILURE_FILL),
    )

    # A compact note keeps the count semantics and cutoff visible without
    # displacing the requested title-first table.
    titles_ws.insert_rows(1, 4)
    titles_ws.merge_cells("A1:G1")
    titles_ws["A1"] = "Role-gate rejection vocabulary audit"
    titles_ws["A1"].font = Font(size=16, bold=True, color="17365D")
    titles_ws["A1"].fill = TITLE_FILL
    titles_ws["A1"].alignment = Alignment(vertical="center")
    titles_ws.row_dimensions[1].height = 28
    titles_ws.merge_cells("A2:G2")
    titles_ws["A2"] = (
        f"Generated {started_at.astimezone(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}; "
        f"central-gate cutoff {cutoff_at.astimezone(timezone.utc).isoformat()}"
    )
    titles_ws.merge_cells("A3:G3")
    titles_ws["A3"] = (
        "Raw occurrences preserve every returned role-rejected row across search instances. "
        "Unique source jobs deduplicate by (source, external ID), matching the production run gate."
    )
    for cell in (titles_ws["A2"], titles_ws["A3"]):
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.font = Font(italic=True, color="404040")
    titles_ws.row_dimensions[3].height = 34
    titles_ws.freeze_panes = "A5"
    # Table's ref must follow the inserted rows.
    for table in titles_ws.tables.values():
        table.ref = f"A4:G{titles_ws.max_row}"
    titles_ws.auto_filter.ref = f"A4:G{titles_ws.max_row}"
    for cell in titles_ws[4]:
        cell.fill = HEADER_FILL
        cell.font = Font(color="FFFFFF", bold=True)

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.row == 1 or cell.row == 4 and ws is titles_ws:
                    continue
                cell.border = Border(bottom=THIN_BLUE) if cell.row == 1 else cell.border
    wb.active = 0
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


def run(*, cutoff: str = DEFAULT_CUTOFF, output_dir: Path | None = None) -> Path:
    started_at = datetime.now(timezone.utc)
    connectors = [*ACTIVE_CONNECTORS, *HEADED_CONNECTORS]
    cutoff_at = parse_since(cutoff)
    apply_sync_window(connectors, since_at=cutoff_at, cutoff_at=started_at)
    enforce_workload_bounds(connectors)
    logger.info("Starting diagnostic-only role audit: %s instances, cutoff=%s", len(connectors), cutoff_at.isoformat())
    results = fetch_all(connectors)
    retry_failed(results)
    raw_rows, unique_rows = _audit_rows(results, cutoff_at)
    output_dir = output_dir or Path("outputs") / f"all_sources_role_rejections_{started_at.strftime('%Y%m%dT%H%M%SZ')}"
    output = output_dir / "all_sources_role_gate_rejections.xlsx"
    _write_workbook(output, raw_rows=raw_rows, unique_rows=unique_rows, results=results, cutoff_at=cutoff_at, started_at=started_at)
    complete = sum(1 for result in results if result.error is None)
    partial = sum(1 for result in results if result.error is not None and result.jobs)
    failed = len(results) - complete - partial
    logger.info("Workbook saved: %s (raw role rows=%s, unique source jobs=%s; complete=%s partial=%s failed=%s)", output, len(raw_rows), len(unique_rows), complete, partial, failed)
    return output


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--cutoff", default=DEFAULT_CUTOFF, help="Exact ISO-8601 gate cutoff.")
    parser.add_argument("--output-dir", type=Path, default=None)
    return parser.parse_args()


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
    args = _parse_args()
    run(cutoff=args.cutoff, output_dir=args.output_dir)
