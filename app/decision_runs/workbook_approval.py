"""Validated company decisions and contact-search choices from a workbook."""
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook


CONTACT_SEARCH_SELECTIONS = {
    "Automatic (posting-derived)": "auto",
    "Data & AI only": "data_ai",
    "Credit Risk only": "credit_risk",
    "Data & AI + Credit Risk": "both",
    "No contact search": "none",
}

COMPANY_APPROVAL_DECISIONS = {"Approved": True, "Rejected": False}

CONTACT_SEARCH_GROUPS = {
    "data_ai": ["data_ai"],
    "credit_risk": ["credit_risk"],
    "both": ["data_ai", "credit_risk"],
    "none": [],
}


@dataclass(frozen=True)
class WorkbookApproval:
    company_approvals: dict[int, bool]
    contact_search_selections: dict[int, str]

def _is_excel_true(value: object) -> bool:
    """Recognize native and LibreOffice-preserved TRUE cells without evaluating formulas."""
    if value is True:
        return True
    if isinstance(value, str):
        return value.strip().upper() in {"TRUE", "=TRUE()"}
    return False


def groups_for_selection(selection: str) -> list[str] | None:
    """Return an explicit group set, or ``None`` for posting-derived auto."""
    if selection == "auto":
        return None
    try:
        return CONTACT_SEARCH_GROUPS[selection]
    except KeyError as exc:
        raise ValueError(f"invalid stored contact search selection: {selection!r}") from exc


def read_contact_search_selections(path: str | Path, run_id: str) -> dict[int, str]:
    """Compatibility reader for the operator group-cutoff approval path."""
    return read_workbook_approval(path, run_id).contact_search_selections


def read_workbook_approval(path: str | Path, run_id: str) -> WorkbookApproval:
    """Read exact company approvals and contact routing from primary rows."""
    workbook = load_workbook(Path(path), data_only=False, read_only=True)
    try:
        _validate_run_id(workbook, run_id)
        approvals: dict[int, bool] = {}
        selections: dict[int, str] = {}
        for group_number in range(1, 5):
            name = f"Group {group_number}"
            if name not in workbook.sheetnames:
                raise ValueError(f"approval workbook is missing {name!r}")
            sheet = workbook[name]
            headers = _headers(sheet, name)
            if sheet.max_row is None:
                sheet.calculate_dimension(force=True)
            for row in range(2, (sheet.max_row or 1) + 1):
                if not _is_excel_true(sheet.cell(row, headers["primary_job"]).value):
                    continue
                company_id_value = sheet.cell(row, headers["company_id"]).value
                if isinstance(company_id_value, bool) or not isinstance(company_id_value, (int, float)):
                    raise ValueError(f"invalid company_id in {name} row {row}")
                company_id = int(company_id_value)
                if company_id != company_id_value:
                    raise ValueError(f"invalid company_id in {name} row {row}")
                approval_label = sheet.cell(row, headers["company_approval"]).value
                approved = COMPANY_APPROVAL_DECISIONS.get(approval_label)
                if approved is None:
                    raise ValueError(
                        f"invalid company approval in {name} row {row}: {approval_label!r}"
                    )
                displayed = sheet.cell(row, headers["contact_search_selection"]).value
                canonical = CONTACT_SEARCH_SELECTIONS.get(displayed)
                if canonical is None:
                    raise ValueError(
                        f"invalid contact search selection in {name} row {row}: {displayed!r}"
                    )
                prior = selections.get(company_id)
                if prior is not None and prior != canonical:
                    raise ValueError(f"conflicting contact search selections for company {company_id}")
                prior_approval = approvals.get(company_id)
                if prior_approval is not None and prior_approval != approved:
                    raise ValueError(f"conflicting approval decisions for company {company_id}")
                approvals[company_id] = approved
                selections[company_id] = canonical
        return WorkbookApproval(approvals, selections)
    finally:
        workbook.close()


def _headers(sheet, sheet_name: str) -> dict[str, int]:
    headers = {cell.value: cell.column for cell in sheet[1] if cell.value}
    required = {
        "company_id", "primary_job", "company_approval",
        "contact_search_selection",
    }
    missing = required - headers.keys()
    if missing:
        raise ValueError(f"{sheet_name} is missing required approval columns: {', '.join(sorted(missing))}")
    return headers


def _validate_run_id(workbook, run_id: str) -> None:
    if "Run summary" not in workbook.sheetnames:
        raise ValueError("approval workbook is missing 'Run summary'")
    summary = workbook["Run summary"]
    headers = _headers_for_summary(summary)
    actual = summary.cell(2, headers["run_id"]).value
    if actual != run_id:
        raise ValueError(f"approval workbook is for run {actual!r}, not {run_id!r}")


def _headers_for_summary(sheet) -> dict[str, int]:
    headers = {cell.value: cell.column for cell in sheet[1] if cell.value}
    if "run_id" not in headers:
        raise ValueError("approval workbook Run summary is missing 'run_id'")
    return headers
