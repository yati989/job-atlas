"""Read-model workbook for one immutable decision run."""
from collections import Counter
from pathlib import Path

import pandas as pd
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.decision_runs.distributions import salary_distribution
from app.decision_runs.workbook_approval import (
    COMPANY_APPROVAL_DECISIONS,
    CONTACT_SEARCH_SELECTIONS,
)
from app.models.orm import DecisionRun, DecisionRunCompany, DecisionRunFinding, DecisionRunJob
from app.reporting.excel import excel_rows


def _frame(rows: list[dict], columns: list[str]) -> pd.DataFrame:
    return pd.DataFrame(excel_rows(rows), columns=columns)


def _default_company_approvals(
    companies: list[DecisionRunCompany],
) -> set[int]:
    """Default every Group 1 company to approved and every other group rejected."""
    return {
        company.company_id
        for company in companies
        if company.group_number == 1
    }


def _write_group(writer, sheet_name: str, rows: list[dict]) -> None:
    columns = ["company_approval", "contact_search_selection", "company_name", "company_rank", "company_id",
               "company_type", "title", "job_link", "posted_at", "primary_job", "job_id",
               "posting_version_id",
               "work_mode", "effective_salary_lpa", "salary_source", "role_family", "qualified_wlb",
               "review_count", "rejected_evidence"]
    _frame(rows, columns).to_excel(writer, sheet_name=sheet_name, index=False)
    worksheet = writer.sheets[sheet_name]
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Font, PatternFill

    header_fill = PatternFill("solid", fgColor="1F4E78")
    approval_fill = PatternFill("solid", fgColor="FFF2CC")
    approved_fill = PatternFill("solid", fgColor="C6EFCE")
    rejected_fill = PatternFill("solid", fgColor="FFC7CE")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    worksheet.row_dimensions[1].height = 32
    worksheet.freeze_panes = "D2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_view.showGridLines = False
    widths = {
        "company_approval": 20,
        "contact_search_selection": 30,
        "company_name": 32,
        "company_rank": 14,
        "company_id": 12,
        "company_type": 16,
        "title": 42,
        "job_link": 48,
        "posted_at": 24,
    }
    for column_name, width in widths.items():
        letter = worksheet.cell(row=1, column=columns.index(column_name) + 1).column_letter
        worksheet.column_dimensions[letter].width = width
    link_column = columns.index("job_link") + 1
    for row_number, row in enumerate(rows, start=2):
        link = row.get("job_link")
        if link:
            cell = worksheet.cell(row=row_number, column=link_column)
            cell.hyperlink = link
            cell.style = "Hyperlink"
    from openpyxl.worksheet.datavalidation import DataValidation
    approval_column = columns.index("company_approval") + 1
    approval_validation = DataValidation(
        type="list", formula1=f'"{",".join(COMPANY_APPROVAL_DECISIONS)}"', allow_blank=False,
        showInputMessage=True, showErrorMessage=True, showDropDown=False,
    )
    approval_validation.promptTitle = "Company approval"
    approval_validation.prompt = "Approved continues this company; Rejected excludes it."
    approval_validation.errorTitle = "Choose Approved or Rejected"
    approval_validation.error = "Use one of the workbook's company approval choices."
    approval_validation.errorStyle = "stop"
    worksheet.add_data_validation(approval_validation)
    selection_column = columns.index("contact_search_selection") + 1
    selection_validation = DataValidation(
        type="list", formula1=f'"{",".join(CONTACT_SEARCH_SELECTIONS)}"', allow_blank=False,
        showInputMessage=True, showErrorMessage=True, showDropDown=False,
    )
    selection_validation.promptTitle = "Contact-search roles"
    selection_validation.prompt = "Choose posting-derived roles or an explicit function override."
    selection_validation.errorTitle = "Choose a listed option"
    selection_validation.error = "Use one of the workbook's contact-search choices."
    selection_validation.errorStyle = "stop"
    worksheet.add_data_validation(selection_validation)
    for row_number, row in enumerate(rows, start=2):
        if row["primary_job"]:
            approval_cell = worksheet.cell(row=row_number, column=approval_column)
            approval_cell.fill = approval_fill
            approval_cell.font = Font(bold=True)
            approval_cell.alignment = Alignment(horizontal="center")
            approval_validation.add(approval_cell)
            selection_validation.add(worksheet.cell(row=row_number, column=selection_column))
    if rows:
        approval_range = f"A2:A{len(rows) + 1}"
        worksheet.conditional_formatting.add(
            approval_range,
            CellIsRule(operator="equal", formula=['"Approved"'], fill=approved_fill),
        )
        worksheet.conditional_formatting.add(
            approval_range,
            CellIsRule(operator="equal", formula=['"Rejected"'], fill=rejected_fill),
        )


def _write_findings(writer, sheet_name: str, rows: list[dict]) -> None:
    columns = ["job_id", "posting_version_id", "company_id", "source", "title", "job_link",
               "reason", "evidence", "parsed_value", "threshold", "comparison"]
    _frame(rows, columns).to_excel(writer, sheet_name=sheet_name, index=False)
    worksheet = writer.sheets[sheet_name]
    link_column = columns.index("job_link") + 1
    for row_number, row in enumerate(rows, start=2):
        link = row.get("job_link")
        if link:
            cell = worksheet.cell(row=row_number, column=link_column)
            cell.hyperlink = link
            cell.style = "Hyperlink"


def build_decision_workbook(session: Session, run_id: str, out_path: str | Path) -> Path:
    """Render only decision snapshot tables—never current jobs or companies."""
    run = session.get(DecisionRun, run_id)
    if run is None:
        raise ValueError(f"unknown run {run_id}")
    jobs = list(session.execute(select(DecisionRunJob).where(DecisionRunJob.run_id == run_id)).scalars())
    company_rows = list(session.execute(select(DecisionRunCompany).where(DecisionRunCompany.run_id == run_id)).scalars())
    companies = {row.company_id: row for row in company_rows}
    findings = list(session.execute(select(DecisionRunFinding).join(DecisionRunJob).where(DecisionRunJob.run_id == run_id)).scalars())
    finding_by_job: dict[int, list[DecisionRunFinding]] = {}
    for finding in findings:
        finding_by_job.setdefault(finding.decision_run_job_id, []).append(finding)

    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        choices = writer.book.create_sheet("_Contact search choices")
        for row_number, label in enumerate(CONTACT_SEARCH_SELECTIONS, start=1):
            choices.cell(row=row_number, column=1).value = label
        choices.sheet_state = "hidden"
        default_approvals = _default_company_approvals(company_rows)
        first_populated_sheet = None
        for group in range(1, 5):
            rows = []
            for job in sorted((j for j in jobs if j.group_number == group), key=lambda j: (companies[j.company_id].rank, j.within_company_rank or 0, j.id)):
                company = companies[job.company_id]
                rows.append({
                    "company_rank": company.rank, "company_name": job.snapshot.get("company_name"),
                    "company_id": job.company_id,
                    "company_type": company.company_type, "job_id": job.job_id,
                    "posting_version_id": job.posting_version_id, "title": job.snapshot.get("title"),
                    "job_link": job.snapshot.get("job_url"),
                    "posted_at": job.snapshot.get("posted_at"), "primary_job": job.is_primary,
                    "company_approval": (
                        "Approved" if job.company_id in default_approvals else "Rejected"
                    ) if job.is_primary else None,
                    "contact_search_selection": (
                        "Automatic (posting-derived)" if job.is_primary else None
                    ),
                    "work_mode": job.work_mode,
                    "effective_salary_lpa": job.effective_salary_lpa, "salary_source": job.salary_source,
                    "role_family": job.role_family, "qualified_wlb": job.qualified_wlb,
                    "review_count": job.review_count, "rejected_evidence": "; ".join(f.evidence or f.reason_code for f in finding_by_job.get(job.id, [])),
                })
            _write_group(writer, f"Group {group}", rows)
            if rows and first_populated_sheet is None:
                first_populated_sheet = f"Group {group}"

        instructions = [
            {"instruction": "Company approval", "detail": "Use the company_approval dropdown in Column A on each company primary row."},
            {"instruction": "Safe default", "detail": "Every Group 1 company defaults to Approved; every company in Groups 2, 3, and 4 defaults to Rejected."},
            {"instruction": "Contact roles", "detail": "Use contact_search_selection to retain posting-derived roles or choose Data & AI, Credit Risk, both, or no contact search."},
            {"instruction": "Approved", "detail": "Includes that exact company in the approved scope."},
            {"instruction": "Rejected", "detail": "Excludes that company from downstream contact search and outreach."},
            {"instruction": "How to submit", "detail": "Edit the Google Sheet linked from the approval email, then reply in that same thread. No attachment is needed."},
        ]
        _frame(instructions, ["instruction", "detail"]).to_excel(writer, sheet_name="Approval instructions", index=False)

        rejected = []
        anomalies = []
        for job in jobs:
            for finding in finding_by_job.get(job.id, []):
                row = {"job_id": job.job_id, "posting_version_id": job.posting_version_id, "company_id": job.company_id,
                       "source": job.snapshot.get("source"), "title": job.snapshot.get("title"),
                       "job_link": job.snapshot.get("job_url"), "reason": finding.reason_code, "evidence": finding.evidence,
                       "parsed_value": finding.parsed_value, "threshold": finding.threshold, "comparison": finding.comparison}
                (anomalies if finding.kind == "anomaly" else rejected).append(row)
        _write_findings(writer, "Rejected jobs", rejected)
        _write_findings(writer, "Anomalies", anomalies)

        distribution_rows = []
        for group in range(1, 5):
            group_jobs = [j for j in jobs if j.group_number == group]
            distribution_rows.append({"group": group, **salary_distribution([j.effective_salary_lpa for j in group_jobs])})
        _frame(distribution_rows, ["group", "known_n", "unknown_n", "p10", "p25", "p50", "mean", "p75", "p90"]).to_excel(writer, sheet_name="Distributions", index=False)

        role_counts = Counter(j.role_family for j in jobs if j.outcome == "eligible")
        salary_counts = Counter(j.salary_source or "unknown" for j in jobs if j.outcome == "eligible")
        wlb_bands = Counter("unknown" if j.qualified_wlb is None else "high" if j.qualified_wlb >= 4 else "medium" if j.qualified_wlb >= 3 else "low" for j in jobs if j.outcome == "eligible")
        company_types = Counter(row.company_type for row in company_rows)
        summary = [{"run_id": run.id, "since_at": run.since_at, "cutoff_at": run.cutoff_at, "state": run.state,
                    "policy_snapshot": str(run.policy_snapshot), "surviving_companies": len(company_rows),
                    "surviving_jobs": sum(j.outcome == "eligible" for j in jobs), "roles": dict(role_counts),
                    "salary_sources": dict(salary_counts), "wlb_bands": dict(wlb_bands),
                    "company_types": dict(company_types), "outcomes": dict(Counter(j.outcome for j in jobs))}]
        _frame(summary, list(summary[0])).to_excel(writer, sheet_name="Run summary", index=False)
        if first_populated_sheet is not None:
            writer.book.active = writer.book.index(writer.book[first_populated_sheet])
    return out
