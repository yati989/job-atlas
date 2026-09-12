"""Import the user's Approved/Declined workbook choices as an exact scope."""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.orm import (
    PublicSelectionScope,
    PublicSelectionScopeItem,
    PublicWorkflowStage,
)
from app.workflows.selection import freeze_selection_scope


_ALLOWED_DECISIONS = {"approved", "declined"}


def apply_workbook_shortlist(
    session: Session, *, run_id: str, scope_id: int, workbook_path: str | Path,
) -> PublicSelectionScope:
    """Validate a complete exported shortlist and freeze its approved rows."""
    source_scope = session.get(PublicSelectionScope, scope_id)
    if source_scope is None or source_scope.run_id != run_id:
        raise ValueError("shortlist source scope does not belong to this run")
    if source_scope.purpose != "selection":
        raise ValueError("workbook shortlist requires a final selection scope")
    expected = dict(session.execute(select(
        PublicSelectionScopeItem.posting_version_id,
        PublicSelectionScopeItem.company_id,
    ).where(PublicSelectionScopeItem.scope_id == scope_id)).all())

    workbook = load_workbook(Path(workbook_path), read_only=True, data_only=False)
    try:
        if "Shortlist" not in workbook.sheetnames:
            raise ValueError("workbook is missing the Shortlist sheet")
        sheet = workbook["Shortlist"]
        headers = {
            str(cell.value).strip(): index
            for index, cell in enumerate(sheet[1])
            if cell.value is not None
        }
        required = {"follow_up_decision", "posting_version_id", "company_id"}
        missing_headers = sorted(required - set(headers))
        if missing_headers:
            raise ValueError(
                "Shortlist sheet is missing columns: " + ", ".join(missing_headers)
            )
        reviewed: dict[int, tuple[int, str]] = {}
        for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            posting_value = row[headers["posting_version_id"]]
            company_value = row[headers["company_id"]]
            decision_value = row[headers["follow_up_decision"]]
            if posting_value is None and company_value is None and decision_value is None:
                continue
            if isinstance(posting_value, bool) or not isinstance(posting_value, int):
                raise ValueError(f"Shortlist row {row_number} has an invalid posting version ID")
            if posting_value in reviewed:
                raise ValueError(f"Shortlist repeats posting version {posting_value}")
            if isinstance(company_value, bool) or not isinstance(company_value, int):
                raise ValueError(f"Shortlist row {row_number} has an invalid company ID")
            decision = str(decision_value or "").strip().lower()
            if decision not in _ALLOWED_DECISIONS:
                raise ValueError(
                    f"Shortlist row {row_number} must be Approved or Declined"
                )
            reviewed[posting_value] = (company_value, decision)
    finally:
        workbook.close()

    if set(reviewed) != set(expected):
        missing = sorted(set(expected) - set(reviewed))
        extra = sorted(set(reviewed) - set(expected))
        raise ValueError(
            f"Shortlist rows changed; missing posting versions {missing}, unexpected {extra}"
        )
    for posting_id, (company_id, _decision) in reviewed.items():
        if company_id != expected[posting_id]:
            raise ValueError(f"Shortlist company changed for posting version {posting_id}")

    approved = tuple(sorted(
        posting_id
        for posting_id, (_company_id, decision) in reviewed.items()
        if decision == "approved"
    ))
    scope = freeze_selection_scope(
        session,
        run_id=run_id,
        mode="manual",
        posting_version_ids=approved,
        company_ids={posting_id: expected[posting_id] for posting_id in approved},
    )
    export_stage = session.scalar(select(PublicWorkflowStage).where(
        PublicWorkflowStage.run_id == run_id,
        PublicWorkflowStage.name == "export",
    ))
    if export_stage is not None:
        export_stage.detail = {
            **(export_stage.detail or {}),
            "scope_id": scope.id,
            "workbook": str(Path(workbook_path)),
        }
        session.flush()
    return scope
