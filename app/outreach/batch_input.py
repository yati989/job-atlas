"""Validated recipient manifests for explicit standalone outreach batches."""
from __future__ import annotations

import csv
import argparse
import io
import json
import re
from pathlib import Path
from typing import Iterable, Mapping

from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, field_validator, model_validator


class OutreachTarget(BaseModel):
    model_config = ConfigDict(frozen=True, extra="forbid")

    recipient_email: str | None = None
    company_name: str | None = None
    to_name: str | None = None
    recipient_title: str | None = None
    company_id: int | None = None
    prospect_id: int | None = None
    matched_job_id: int | None = None
    resume_path: str | None = None

    @field_validator("recipient_email")
    @classmethod
    def valid_email(cls, value: str | None) -> str | None:
        if value is None or not value.strip():
            return None
        normalized = value.strip().lower()
        if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", normalized):
            raise ValueError("recipient_email is invalid")
        return normalized

    @field_validator("company_name", "to_name", "recipient_title", "resume_path")
    @classmethod
    def clean_optional_text(cls, value: str | None) -> str | None:
        return value.strip() if value and value.strip() else None

    @model_validator(mode="after")
    def address_or_discovery_identity(self) -> "OutreachTarget":
        if self.recipient_email is None and not (self.company_name or self.company_id or self.prospect_id):
            raise ValueError("a recipient email or company identity is required")
        return self

    @property
    def requires_email_discovery(self) -> bool:
        return self.recipient_email is None


_ALIASES = {
    "email": "recipient_email",
    "to_email": "recipient_email",
    "company": "company_name",
    "name": "to_name",
    "title": "recipient_title",
    "job_id": "matched_job_id",
}
_FIELDS = set(OutreachTarget.model_fields)


def _normalize_row(row: Mapping[str, object], row_number: int) -> OutreachTarget:
    normalized: dict[str, object] = {}
    for key, value in row.items():
        name = _ALIASES.get(str(key).strip().lower(), str(key).strip().lower())
        if name not in _FIELDS:
            continue
        if value == "":
            value = None
        normalized[name] = value
    try:
        return OutreachTarget.model_validate(normalized)
    except Exception as exc:
        raise ValueError(f"invalid outreach target at row {row_number}: {exc}") from exc


def _targets(rows: Iterable[Mapping[str, object]], *, first_data_row: int) -> tuple[OutreachTarget, ...]:
    return tuple(
        _normalize_row(row, row_number)
        for row_number, row in enumerate(rows, start=first_data_row)
    )


def parse_pasted_targets(text: str) -> tuple[OutreachTarget, ...]:
    """Parse a pasted CSV table or one email address per line."""
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if not lines:
        raise ValueError("recipient list is empty")
    if "," not in lines[0] and "@" in lines[0]:
        return _targets(
            ({"recipient_email": line} for line in lines), first_data_row=1
        )
    return _targets(csv.DictReader(io.StringIO("\n".join(lines))), first_data_row=2)


def load_outreach_targets(path: str | Path) -> tuple[OutreachTarget, ...]:
    """Load CSV, JSON, or XLSX targets through one validated row contract."""
    path = Path(path).expanduser()
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as handle:
            return _targets(csv.DictReader(handle), first_data_row=2)
    if suffix == ".json":
        data = json.loads(path.read_text(encoding="utf-8"))
        if not isinstance(data, list) or not all(isinstance(row, Mapping) for row in data):
            raise ValueError("JSON outreach input must be a list of row objects")
        return _targets(data, first_data_row=1)
    if suffix == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            values = sheet.iter_rows(values_only=True)
            headers = next(values, None)
            if not headers:
                raise ValueError("Excel outreach input is empty")
            names = [str(value or "").strip() for value in headers]
            rows = (dict(zip(names, values_row)) for values_row in values)
            return _targets(rows, first_data_row=2)
        finally:
            workbook.close()
    raise ValueError("outreach input must be .csv, .json, or .xlsx")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Validate a standalone outreach CSV, JSON, or Excel batch."
    )
    inputs = parser.add_mutually_exclusive_group(required=True)
    inputs.add_argument("--input", type=Path)
    inputs.add_argument("--paste")
    args = parser.parse_args(argv)
    targets = (
        load_outreach_targets(args.input)
        if args.input is not None
        else parse_pasted_targets(args.paste)
    )
    payload = {
        "target_count": len(targets),
        "requires_email_discovery": sum(
            target.requires_email_discovery for target in targets
        ),
        "targets": [target.model_dump(mode="json") for target in targets],
    }
    print(json.dumps(payload, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
